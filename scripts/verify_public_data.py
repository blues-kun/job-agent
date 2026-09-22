"""核对公开数据的摘要、行数、双格式一致性和当前清理规则；仅输出聚合结果。"""
import argparse
import csv
import gzip
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from scripts.export_public_jobs import ARTIFACTS, HEADERS, PATTERNS, actionable_matches


def verify(directory: Path) -> dict:
    manifest = json.loads((directory / ARTIFACTS[2]).read_text(encoding="utf-8"))
    for name in ARTIFACTS[:2]:
        path = directory / name
        expected = manifest["files"][name]
        if path.stat().st_size != expected["bytes"] or hashlib.sha256(path.read_bytes()).hexdigest() != expected["sha256"]:
            raise ValueError("公开文件摘要不匹配：" + name)
    with gzip.open(directory / ARTIFACTS[1], "rt", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != HEADERS: raise ValueError("CSV字段不匹配")
        rows = list(reader)
    book = load_workbook(directory / ARTIFACTS[0], read_only=True, data_only=False)
    try:
        if len(book.worksheets) != 1: raise ValueError("公开表出现额外工作表")
        iterator = book.active.iter_rows()
        if [cell.value for cell in next(iterator)] != HEADERS: raise ValueError("Excel字段不匹配")
        i = 0
        for cells in iterator:
            if any(cell.data_type == "f" for cell in cells): raise ValueError("公开表含公式")
            record = {key: cell.value or "" for key, cell in zip(HEADERS, cells)}
            if i >= len(rows) or record != rows[i]: raise ValueError("XLSX与CSV内容不一致")
            i += 1
        if i != len(rows) or i != manifest["rows"]: raise ValueError("行数不匹配")
    finally:
        book.close()
    logical_hash = hashlib.sha256()
    for row in rows:
        if any(actionable_matches(kind, value) for value in row.values() for kind in PATTERNS):
            raise ValueError("公开内容命中联系方式规则，需单独复核")
        logical_hash.update((json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode())
    if logical_hash.hexdigest() != manifest["logical_content_sha256"]: raise ValueError("逻辑内容摘要不匹配")
    if len({row["记录编号"] for row in rows}) != len(rows): raise ValueError("记录编号重复")
    companies = len({row["企业"] for row in rows})
    keys = len({tuple(row[k] for k in ["岗位名称", "企业", "岗位薪资"]) for row in rows})
    if companies != manifest["company_count"] or keys != manifest["business_key_unique"]:
        raise ValueError("企业/业务键统计不一致")
    return {"rows": len(rows), "business_key_unique": keys, "company_count": companies,
            "formats_equal": True, "hashes_verified": True, "contact_patterns_clear": True,
            "note": "规则复核通过不构成无法再识别或所有个人信息已覆盖的保证。"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", type=Path, default=Path("data"))
    args = parser.parse_args()
    print(json.dumps(verify(args.directory), ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
