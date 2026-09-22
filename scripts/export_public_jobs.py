"""只读原表，导出经过信息清理的公开岗位 XLSX、压缩 CSV 与来源清单。"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import csv
import gzip
import hashlib
import io
import json
import os
from pathlib import Path
import re
import tempfile
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell


VERSION = "public-jobs-v1"
FIELDS = ["岗位名称", "企业", "岗位薪资", "岗位要求", "岗位职责", "岗位地址", "职位类型名称", "二级分类"]
HEADERS = [*FIELDS, "记录编号"]
ARTIFACTS = ["job_data_public.xlsx", "job_data_public.csv.gz", "job_data_public.manifest.json"]
MARKER = "[联系信息已移除]"
# 联系方式句段整体移除，技术名词“微信小程序”“QQ小程序”不单独触发。
CONTACT_CLAUSE = re.compile(
    r"(?:联系人|联系方式|联系电话|联系手机|手机号|电话[：:]|手机[：:]|邮箱|电子邮件|"
    r"请联系|联系我|联系[：:]|投递[至到]|发送[至到]|微信号|微信[：:]|加我微信|加微信|"
    r"(?:wechat|weixin|vx|v信|薇信|QQ号?|扣扣|企鹅号)\s*[：:=]|(?:工作|办公|面试|公司|详细|上班)地[址点]\s*[：:]|"
    r"[\u4e00-\u9fff]{1,3}(?:先生|女士|小姐)|(?:联系|咨询|添加|致电|找)[^\n。；;]{0,12}?[\u4e00-\u9fff]{1,3}老师|"
    r"(?:招聘联系人|招聘负责人|招聘专员|招聘顾问|简历接收人|联络人|对接人|招聘|人事|HR)\s*[：:]\s*[\u4e00-\u9fffA-Za-z]{2,})", re.I)
TECHNICAL_NAMES = {"asp.net", "socket.io", "vb.net"}
PATTERNS = {
    "网址": re.compile(r"(?:https?://|ftp://|www\.)[^\s<>，。；;]+", re.I),
    "邮箱": re.compile(r"[A-Za-z0-9_.+%-]+\s*@\s*[A-Za-z0-9.-]+\s*\.\s*[A-Za-z]{2,}"),
    "混淆邮箱": re.compile(r"[A-Za-z0-9_.+-]+\s*(?:\[\s*at\s*\]|\(\s*at\s*\)|艾特|\s+at\s+)\s*[A-Za-z0-9-]+(?:\s*(?:\.|\[\s*dot\s*\]|\(\s*dot\s*\)|点)\s*[A-Za-z0-9-]+)+", re.I),
    "裸域名": re.compile(r"(?<![\w@.])[A-Za-z0-9][A-Za-z0-9-]{0,62}(?:\.[A-Za-z0-9-]{1,63})*\.(?:com|cn|net|org|io|me|app|work|tech|co|xyz|link|top|cc|vip|ltd|site|info)(?:/[^\s<>，。；;]*)?(?![\w.])", re.I),
    "社交账号": re.compile(r"(?:微[信芯]|薇[信芯]|weixin|wechat|v\s*x|v信|加v)[ \t]*(?:(?:账号|帐号|号码|号|是|为|联系|咨询|添加|同号|[:=+\-])[ \t]*){0,5}[A-Za-z0-9][A-Za-z0-9_.\-]{4,29}", re.I),
    "QQ群号": re.compile(r"(?<![A-Za-z])(?:Q\s*Q(?:群)?|Q群|群号|扣扣|企鹅号)[ \t]*(?:号|号码|[:=+])?[ \t]*[1-9][0-9 \t\-]{4,16}\d(?!\d)", re.I),
    "手机": re.compile(r"(?<!\d)(?:(?:\+?86)[ -]*)?1[3-9](?:[ \t().-]*\d){9}(?!\d)"),
    "座机": re.compile(r"(?<!\d)(?:\(0\d{2,3}\)|0\d{2,3})[ \t-]*(?:\d[ \t-]?){6,7}\d(?!\d)"),
    "国际电话": re.compile(r"(?<!\d)\+(?:852|853|886|1|44|65|81|82)[ \t-]*(?:\d[ \t-]*){7,12}(?!\d)"),
    "中文号码": re.compile(r"(?<![零〇一二两三四五六七八九幺])(?:一|幺)[三四五六七八九][零〇一二两三四五六七八九幺]{9}(?![零〇一二两三四五六七八九幺])"),
    "长号码": re.compile(r"(?<!\d)\d{8,}[Xx]?(?!\d)"),
    "凭据": re.compile(r"(?:sk-|ghp_|github_pat_)[A-Za-z0-9_-]{16,}"),
}
CITIES = ("深圳", "广州", "北京", "上海", "杭州", "东莞", "惠州", "佛山", "珠海", "中山", "苏州", "南京", "成都", "武汉", "西安", "重庆", "长沙", "合肥", "厦门", "天津", "青岛", "香港")
DISTRICTS = ("南山区", "福田区", "宝安区", "龙岗区", "龙华区", "罗湖区", "光明区", "坪山区", "盐田区", "大鹏新区", "深汕特别合作区")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalized(value) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    # 去除零宽格式字符与非法控制字符，保留段落换行。
    return "".join(c for c in text if c in "\n\t" or unicodedata.category(c) not in {"Cf", "Cc"}).strip()


def actionable_matches(kind: str, text: str):
    return [m for m in PATTERNS[kind].finditer(text)
            if not (kind == "裸域名" and m.group().lower() in TECHNICAL_NAMES)]


def sanitize_text(value, company: str, counts: Counter) -> str:
    text = normalized(value)
    # 同名企业只替换完整标识，技术产品名称和其他公开实体不作实体匿名承诺。
    if len(company) >= 2:
        occurrences = text.count(company)
        if occurrences:
            counts["正文同名企业"] += occurrences
            text = text.replace(company, "[企业名称已编码]")
    clauses = re.split(r"([\n。；;])", text)
    for index in range(0, len(clauses), 2):
        if CONTACT_CLAUSE.search(clauses[index]):
            clauses[index] = MARKER
            counts["联系或详细地址句段"] += 1
    text = "".join(clauses)
    for kind, pattern in PATTERNS.items():
        matches = actionable_matches(kind, text)
        for match in reversed(matches):
            text = text[:match.start()] + MARKER + text[match.end():]
        counts[kind] += len(matches)
    # CSV 与 XLSX 都使用相同安全文本，防止下载后被表格软件执行为公式。
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
        counts["表格公式前缀"] += 1
    return text


def coarse_address(value) -> str:
    text = normalized(value)
    city = next((city for city in CITIES if city in text), "")
    district = next((district for district in DISTRICTS if district in text), "") if city in {"", "深圳"} else ""
    if city and not district:
        found = re.search(re.escape(city) + r"(?:市)?([\u4e00-\u9fff]{2,4}?[区县])", text)
        district = found.group(1) if found else ""
    return city + ("市" if city and city != "香港" else "") + district


def read_source(source: Path) -> tuple[list[dict], int]:
    rows = []
    book = load_workbook(source, read_only=True, data_only=False)
    try:
        sheets = len(book.worksheets)
        for sheet in book:
            iterator = sheet.iter_rows(values_only=True)
            header = [str(value or "").strip().lstrip("\ufeff") for value in next(iterator, [])]
            if header != FIELDS:
                raise ValueError("原表须为已审核的8列模式；新增字段需单独审查，不能直接公开")
            for values in iterator:
                if any(value is not None for value in values):
                    rows.append({key: str(value).strip() if value is not None else "" for key, value in zip(header, values)})
    finally:
        book.close()
    if not rows or any(not row["岗位名称"] or not row["企业"] for row in rows):
        raise ValueError("原表为空或缺岗位名称/企业，需先核对数据")
    return rows, sheets


def export_public_jobs(source: Path, output: Path) -> dict:
    source, output = Path(source).resolve(), Path(output).absolute()
    if any(p.is_symlink() for p in [output, *output.parents]):
        raise ValueError("输出目录不能经过符号链接")
    if any((output / name).exists() or (output / name).is_symlink() for name in ARTIFACTS):
        raise FileExistsError("公开产物已存在；请使用新的导出目录")
    before = sha256(source)
    rows, sheets = read_source(source)
    companies = {name: f"企业{index:05d}" for index, name in enumerate(sorted({row["企业"] for row in rows}), 1)}
    public, substitutions, changed_fields = [], Counter(), Counter()
    for index, row in enumerate(rows, 1):
        item = {}
        for field in FIELDS:
            if field == "企业":
                value = companies[row[field]]
            elif field == "岗位地址":
                value = coarse_address(row[field])
            else:
                value = sanitize_text(row[field], normalized(row["企业"]), substitutions)
            if value != row[field]:
                changed_fields[field] += 1
            item[field] = value
        item["记录编号"] = f"J{index:06d}"
        public.append(item)
    # 号码、邮件和URL规则必须在输出文本上再次通过，公开清单仅保存计数。
    residual = {kind: sum(bool(actionable_matches(kind, value)) for row in public for value in row.values()) for kind in PATTERNS}
    if any(residual.values()):
        raise ValueError("导出内容仍命中联系方式规则，停止生成公开文件")
    if sha256(source) != before:
        raise ValueError("读取期间原表发生变化，停止导出")
    output.mkdir(parents=True, exist_ok=True)
    logical_hash = hashlib.sha256()
    for row in public:
        logical_hash.update((json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode())
    manifest = {
        "version": VERSION, "created_at": datetime.now(timezone.utc).isoformat(),
        "source_sha256": before, "source_sheets": sheets, "source_rows": len(rows), "rows": len(public),
        "columns": HEADERS, "company_count": len(companies),
        "category_count": len({row["职位类型名称"] for row in public}),
        "category_parent_count": len({row["二级分类"] for row in public}),
        "business_key_unique": len({tuple(row[key] for key in ["岗位名称", "企业", "岗位薪资"]) for row in public}),
        "logical_content_sha256": logical_hash.hexdigest(), "exporter_sha256": sha256(Path(__file__)),
        "changed_cells_by_column": dict(changed_fields), "substitutions": dict(substitutions),
        "residual_pattern_cells": residual,
        "notes": ["来自历史岗位广告，采集日期与招聘有效状态未知。",
                  "企业为本版本稳定编号，未发布企业映射；保留技术与产品名称。",
                  "地址仅保留可识别城市/行政区；无法识别保持空值。",
                  "自动规则清理不构成无法再识别保证；不可用来恢复个人联系方式。",
                  "公开版有独立内容摘要，原实验标签、引用跨度和模型缓存不可直接套用。"],
        "files": {},
    }
    with tempfile.TemporaryDirectory(prefix=".public-export-", dir=output) as temporary:
        staging = Path(temporary)
        book = Workbook(write_only=True)
        sheet = book.create_sheet("岗位公开版")
        for values in [HEADERS, *[[row[key] for key in HEADERS] for row in public]]:
            cells = []
            for value in values:
                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
        book.save(staging / ARTIFACTS[0]); book.close()
        with (staging / ARTIFACTS[1]).open("wb") as raw:
            with gzip.GzipFile(filename="", mode="wb", fileobj=raw, mtime=0) as zipped:
                with io.TextIOWrapper(zipped, encoding="utf-8-sig", newline="") as stream:
                    writer = csv.DictWriter(stream, fieldnames=HEADERS)
                    writer.writeheader(); writer.writerows(public)
        for name in ARTIFACTS[:2]:
            manifest["files"][name] = {"sha256": sha256(staging / name), "bytes": (staging / name).stat().st_size}
        (staging / ARTIFACTS[2]).write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if sha256(source) != before:
            raise ValueError("导出期间原表发生变化，停止发布")
        for name in ARTIFACTS:
            # 硬链接提供禁止覆盖语义，产物来自同一输出目录中的临时文件。
            os.link(staging / name, output / name)
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True, help="只读原始岗位xlsx路径")
    parser.add_argument("--output-dir", type=Path, required=True, help="未包含公开产物的导出目录")
    args = parser.parse_args()
    result = export_public_jobs(args.source, args.output_dir)
    print(json.dumps({key: result[key] for key in ["version", "rows", "company_count", "category_count", "business_key_unique", "changed_cells_by_column", "residual_pattern_cells"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
