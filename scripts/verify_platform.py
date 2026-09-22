"""对本机已启动平台做可复跑验收，只输出虚构样例的聚合结果。"""
from __future__ import annotations

import argparse
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import time
from urllib.parse import urlparse

import httpx
from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://127.0.0.1:8090")
    parser.add_argument("--source", type=Path, default=Path(__file__).resolve().parents[1]/"data/job_data.xlsx")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--dataset",type=Path,help="版本化岗位目录；不用旧业务键定位证据")
    args = parser.parse_args()
    if urlparse(args.url).hostname not in {"localhost", "127.0.0.1", "::1"}:
        parser.error("验收脚本只访问本机服务")
    if args.output.resolve() == args.source.resolve():
        parser.error("结果文件不能覆盖原始岗位库")
    before = hashlib.sha256(args.source.read_bytes()).hexdigest()
    # 不重建检索索引；直接从原表定位引用，独立核对字段与偏移。
    records = {}
    book = load_workbook(args.source, read_only=True, data_only=True)
    try:
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = [str(value or "").strip().strip("\ufeff") for value in next(rows, [])]
            for row in rows:
                data = {key:str(value).strip() if value is not None else "" for key,value in zip(header,row)}
                def get(*keys):
                    return next((data[key] for key in keys if data.get(key)), "")
                key = (get("岗位名称","职位名称","jobName"), get("企业","公司","公司名称","brandName"), get("岗位薪资","薪资","salaryDesc"))
                job_id = hashlib.sha256("\0".join(key).encode()).hexdigest()[:20]
                records.setdefault(job_id, {"description":get("岗位职责","职位描述","岗位描述"), "requirements":get("岗位要求","任职要求","职位要求")})
    finally:
        book.close()
    if args.dataset:
        manifest=json.loads((args.dataset/"manifest.json").read_text())
        content=(args.dataset/"jobs.jsonl").read_bytes()
        assert manifest["source_sha256"]==before
        assert hashlib.sha256(content).hexdigest()==manifest["files"]["jobs.jsonl"]["sha256"]
        records={row["job_id"]:row for row in (json.loads(line) for line in content.decode().splitlines() if line)}
    client = httpx.Client(base_url=args.url.rstrip("/"), timeout=60, trust_env=False)
    def call(path, payload=None, method=None):
        response=client.request(method or ("GET" if payload is None else "POST"), path, json=payload)
        response.raise_for_status()
        return response.json()
    def confirmed_payload(payload):
        preview=call("/api/v2/profile/preview",payload)
        assert not preview["conflicts"], "虚构样例画像与其表单存在冲突，请先修正验收材料"
        current={**payload,"preferences":preview["preferences"],
                 "field_origins":{field["key"]:field["source"] for field in preview["fields"]}}
        result=call("/api/v2/profile/confirm",{**current,"user_confirmed":True,"acknowledged_conflicts":[]})
        assert result["preferences"]==current["preferences"]
        return {**current,"confirmation_token":result["confirmation_token"]}
    result = {"executed_at_utc":datetime.now(timezone.utc).isoformat(), "source_sha256":before, "samples":[],
              "scope":"虚构演示样例的工程验收；不代表独立人工人岗效果", "human_gold_created":False}
    confirmation_checks=0
    expected_feature_version="evidence-ranker-v2-35"
    quote_count = 0
    hard_violations = 0
    try:
        result["health"] = call("/api/v2/health")
        assert result["health"]["snapshot"] == before[:16], "服务快照与本地验收文件不同"
        overview = call("/api/v2/overview")
        result["engine"] = overview["engine"]
        result["overview"] = {key:overview[key] for key in ["raw_total","total","duplicates","category_count","unknown_location","salary_count","salary_quantiles","alternative_groups"]}
        samples = call("/api/v2/samples")["samples"]
        selected = None
        for sample in samples:
            payload = {"text":sample["text"],"preferences":sample["preferences"],
                       "field_origins":{key:"sample" for key in sample["preferences"]}}
            assert client.post("/api/v2/recommend",json=payload).status_code==409
            payload=confirmed_payload(payload)
            confirmation_checks+=1
            output = call("/api/v2/recommend",payload)
            assert output["feature_version"]==expected_feature_version
            for job in output["jobs"]:
                assert len(job["features"])==35
                assert {"task_alignment","requirement_support_ratio","requirement_unknown_ratio"}.issubset(job["features"])
                assert job["id"] in records
                hard_violations += sum(item["status"]=="fail" for item in job["constraints"])
                for match in job["matched"]:
                    jd, resume = match["job_evidence"], match["resume_evidence"]
                    assert records[job["id"]][jd["field"]][jd["start"]:jd["end"]] == jd["quote"]
                    assert sample["text"][resume["start"]:resume["end"]] == resume["quote"]
                    quote_count += 1
            expected = "clarify" if sample["id"] in {"incomplete","keyword-list"} else "recommend"
            assert output["action"] == expected
            result["samples"].append({"sample":sample["name"],"id":sample["id"],"action":output["action"],
                "returned":len(output["jobs"]),"eligible":output["total_eligible"],"latency_ms":output["latency_ms"],
                "sources":output.get("source_counts",{}),"retriever":output["retriever"],
                "avg_skill_coverage":round(sum(job["skill_coverage"] for job in output["jobs"])/len(output["jobs"]),2) if output["jobs"] else None})
            if sample["id"] == "python-junior":
                selected = payload, output
        assert hard_violations == 0
        assert selected is not None
        payload, output = selected
        assert output["jobs"], "Python虚构验收样例未返回岗位"
        # token绑定原文、偏好和会话，任何后续编辑都必须重新确认。
        assert client.post("/api/v2/recommend",json={**payload,"text":payload["text"]+"\n补充另一条待确认的经历。"}).status_code==409
        changed_preferences={**payload["preferences"],"experience_years":payload["preferences"]["experience_years"]+1}
        assert client.post("/api/v2/recommend",json={**payload,"preferences":changed_preferences}).status_code==409
        confirmation_checks+=2
        target = {**payload,"job_id":output["jobs"][0]["id"]}
        diagnosis = call("/api/v2/diagnose", target)
        rewrite = call("/api/v2/rewrite", target)
        interview = call("/api/v2/interview", target)
        assert diagnosis["job"]["id"] == target["job_id"] and interview["questions"]
        assert rewrite["facts_preserved"] and rewrite["new_claims"] == 0
        assert Counter(line.strip() for line in rewrite["revised"].splitlines() if line.strip()) == Counter(line.strip() for line in payload["text"].splitlines() if line.strip())
        assert rewrite["changed_fields"]=={} and rewrite["added_blocks"]==[] and rewrite["removed_blocks"]==[]
        assert rewrite["added_skill_statements"]==[]
        feedback = {"run_id":output["run_id"],"job_id":target["job_id"],"action":"like"}
        call("/api/v2/feedback",feedback)
        call("/api/v2/feedback",feedback)
        assert call("/api/v2/feedback")["counts"] == {"like":1}
        assert call("/api/v2/research-events")["events"]==[], "默认不应留存研究特征"
        consented=call("/api/v2/recommend",{**payload,"research_consent":True})
        events=call("/api/v2/research-events")["events"]
        assert len(events)==1 and events[0]["raw_resume_stored"] is False and events[0]["consent"] is True
        assert events[0]["profile_version"]==consented["profile"]["version"]
        assert events[0]["feature_version"]==expected_feature_version
        assert all(len(row["features"])==35 for row in events[0]["exposures"])
        assert payload["text"] not in json.dumps(events,ensure_ascii=False)
        result["comparison"] = call("/api/v2/compare",payload)
        sequential = []
        for _ in range(3):
            started=time.perf_counter();call("/api/v2/recommend",payload)
            sequential.append(round((time.perf_counter()-started)*1000,1))
        def concurrent_run(_):
            started=time.perf_counter();call("/api/v2/recommend",payload)
            return round((time.perf_counter()-started)*1000,1)
        result["sequential_http_ms"] = sequential
        with ThreadPoolExecutor(max_workers=5) as pool:
            result["concurrent5_http_ms"] = list(pool.map(concurrent_run, range(5)))
        for path in ["/","/assets/styles.css","/assets/app.js"]:
            assert client.get(path).status_code == 200
        assert client.get("/data/job_data.xlsx").status_code == 404
        assert client.post("/api/v2/recommend",json=payload,headers={"Origin":"https://untrusted.invalid"}).status_code == 403
        result["checks"] = {"sample_runs":len(samples),"hard_violations_detected_by_current_rules":hard_violations,
            "verified_dual_quotes":quote_count,"diagnose_rewrite_interview":True,"feedback_idempotent":True,
            "explicit_profile_confirmation":True,"confirmation_gate_checks":confirmation_checks,
            "feature_version":expected_feature_version,"feature_count":35,"rewrite_fields_and_blocks_preserved":True,
            "research_features_require_consent":True,"research_events_exclude_raw_resume":True,
            "static_assets":True,"source_hash_unchanged":hashlib.sha256(args.source.read_bytes()).hexdigest()==before}
        assert result["checks"]["source_hash_unchanged"]
    finally:
        # 本客户端使用独立会话；清理验收曝光，避免把虚构样例当成训练反馈。
        try:
            try:
                call("/api/v2/feedback", method="DELETE")
            finally:
                call("/api/v2/journey", method="DELETE")
                assert call("/api/v2/research-events")["events"]==[]
        finally:
            client.close()
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"result_file":str(args.output),"checks":result["checks"]},ensure_ascii=False))


if __name__ == "__main__":
    main()
