"""简历文件内存解析；不保存上传原件，不进行OCR或执行嵌入内容。"""
from io import BytesIO
from pathlib import Path
import zipfile


def extract_document(content:bytes,filename:str):
    if not content or len(content)>2*1024*1024:raise ValueError("文件需为2MB以内的非空TXT、PDF或DOCX")
    suffix=Path(filename).suffix.lower()
    pages=[]
    if suffix==".txt":
        pages=[{"page":None,"text":content.decode("utf-8-sig")}]
    elif suffix==".pdf":
        from pypdf import PdfReader
        reader=PdfReader(BytesIO(content),strict=True)
        if reader.is_encrypted:raise ValueError("请先解除PDF密码")
        if len(reader.pages)>20:raise ValueError("简历最多20页")
        for number,page in enumerate(reader.pages,1):
            stream=page.get_contents()
            if stream is not None and len(stream.get_data())>8*1024*1024:raise ValueError("PDF页面过大")
            pages.append({"page":number,"text":page.extract_text() or ""})
    elif suffix==".docx":
        from docx import Document
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries=archive.infolist()
            if len(entries)>1000 or sum(item.file_size for item in entries)>20*1024*1024:raise ValueError("DOCX解压后过大")
        document=Document(BytesIO(content))
        chunks=[paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:chunks.append(" | ".join(cell.text for cell in row.cells))
        pages=[{"page":None,"text":"\n".join(chunks)}]
    else:raise ValueError("仅支持UTF-8 TXT、文字型PDF和DOCX")
    text="\n".join(item["text"] for item in pages).strip()
    if not text:raise ValueError("没有提取到文字；扫描件请先转成可选择文字的PDF")
    if len(text)>20000:raise ValueError("提取文本超过20,000字，请保留相关经历")
    return {"text":text,"pages":pages,"note":"文件仅在内存解析；请核对阅读顺序、日期和表格内容，扫描件未执行OCR。"}
