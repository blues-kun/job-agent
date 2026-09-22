#!/usr/bin/env python3
"""只读招聘 XLSX，生成仓库外可追溯、按岗位家族隔离的研究数据。

本工具只生成机器解析与待标注任务，不生成或冒充人工标签。
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict, deque
from contextlib import contextmanager
from datetime import datetime, timezone
import hashlib
from itertools import combinations
import json
import math
import os
from pathlib import Path
import random
import re
import shutil
import sys
import tempfile
from typing import Any

SCHEMA_VERSION = "job-agent-research-data-v1"
SPLITS = ("train", "dev", "test")
FIELDS = {
    "title": ("岗位名称", "职位名称", "职位名", "jobName", "title"),
    "company": ("企业", "公司", "公司名称", "企业名称", "brandName", "company"),
    "salary_raw": ("岗位薪资", "薪资", "salaryDesc", "salary"),
    "category": ("职位类型名称", "三级分类", "职位类型", "岗位类别", "category"),
    "category_parent": ("二级分类", "大类", "一级分类", "category_parent"),
    "requirements": ("岗位要求", "任职要求", "职位要求", "requirements"),
    "description": ("岗位职责", "职位描述", "岗位描述", "description"),
    "address": ("岗位地址", "工作地点", "工作地址", "地址", "address"),
}


def canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def value_hash(value: Any) -> str:
    return hashlib.sha256(canonical(value).encode("utf-8")).hexdigest()


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized(text: str) -> str:
    return re.sub(r"\s+", "", text).casefold()


def load_domain(repository: Path):
    if not (repository / "job_agent/domain.py").is_file():
        raise ValueError("--repo-root 中未找到 job_agent/domain.py。")
    sys.path.insert(0, str(repository))
    from job_agent.domain import Job
    return Job


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def read_jobs(source: Path, Job, snapshot: str) -> tuple[list[dict], dict]:
    """内容哈希覆盖所有原始列；业务键相同但职责不同的记录不会丢失。"""
    from openpyxl import load_workbook
    by_hash: dict[str, dict] = {}
    short_hashes: dict[str, str] = {}
    sheets, raw_count, invalid_title = [], 0, 0
    book = load_workbook(source, read_only=True, data_only=False)
    try:
        for sheet_index, sheet in enumerate(book.worksheets, 1):
            iterator = sheet.iter_rows(values_only=True)
            raw_header = next(iterator, ())
            seen = Counter()
            headers, header_base = [], {}
            for index, cell in enumerate(raw_header, 1):
                base = cell_text(cell).strip().strip("\ufeff").strip() or f"未命名列{index}"
                seen[base] += 1
                header = base if seen[base] == 1 else f"{base}__重复{seen[base]}"
                headers.append(header)
                header_base[header] = base
            count = 0
            for row_number, row in enumerate(iterator, 2):
                if not any(value is not None and cell_text(value).strip() for value in row):
                    continue
                raw_count += 1
                count += 1
                raw = {header: cell_text(value) for header, value in zip(headers, row)}
                # 保留完整列内容，包括无已知映射的列；只在业务视图中修剪首尾空白。
                full_hash = value_hash(raw)
                job_id = full_hash[:20]
                if job_id in short_hashes and short_hashes[job_id] != full_hash:
                    raise RuntimeError("20位内容ID发生哈希碰撞，拒绝生成数据。")
                short_hashes[job_id] = full_hash
                source_id = f"sha256:{snapshot}/sheet:{sheet_index}/row:{row_number}"
                if full_hash in by_hash:
                    by_hash[full_hash]["source_record_ids"].append(source_id)
                    continue
                values, sources = {}, {}
                for field, aliases in FIELDS.items():
                    candidates = [h for alias in aliases for h in headers if header_base[h] == alias]
                    selected = next((h for h in candidates if raw.get(h, "").strip()), candidates[0] if candidates else None)
                    values[field] = raw.get(selected, "").strip() if selected else ""
                    sources[field] = selected
                legacy_id = hashlib.sha256("\0".join(values[field] for field in ("title", "company", "salary_raw")).encode("utf-8")).hexdigest()[:20]
                parsed = Job(id=job_id, version=job_id, title=values["title"], company=values["company"], category=values["category"],
                             family=values["category_parent"], salary_raw=values["salary_raw"], requirements=values["requirements"],
                             description=values["description"], address=values["address"], row=row_number, sheet=sheet_index)
                missing_title = not bool(values["title"])
                invalid_title += missing_title
                by_hash[full_hash] = {
                    "job_id": job_id, "job_version_id": job_id, "content_sha256": full_hash, "legacy_id": legacy_id,
                    "source_record_ids": [source_id], "snapshot": snapshot, **values,
                    "company_hash": hashlib.sha256(normalized(values["company"]).encode("utf-8")).hexdigest()[:20],
                    "source_fields": raw, "field_sources": sources,
                    "skills": parsed.skills, "groups": parsed.groups, "salary": parsed.salary,
                    "quality_flags": ["missing_title"] if missing_title else [],
                }
            sheets.append({"sheet_number": sheet_index, "raw_rows": count, "column_count": len(headers)})
    finally:
        book.close()
    jobs = sorted(by_hash.values(), key=lambda row: row["job_id"])
    legacy = Counter(job["legacy_id"] for job in jobs)
    return jobs, {"raw_records": raw_count, "content_unique_records": len(jobs), "content_duplicate_records": raw_count-len(jobs),
                  "legacy_key_collision_groups": sum(count > 1 for count in legacy.values()),
                  "legacy_key_additional_versions": sum(count-1 for count in legacy.values()),
                  "missing_title_records": invalid_title, "sheets": sheets}


def assign_families(jobs: list[dict], seed: int, ratios: tuple[float, float, float]) -> dict:
    parents = list(range(len(jobs)))
    def find(index):
        while parents[index] != index:
            parents[index] = parents[parents[index]]
            index = parents[index]
        return index
    def union(left, right):
        left, right = find(left), find(right)
        if left != right:
            parents[max(left, right)] = min(left, right)
    seen_description, seen_company_title = {}, {}
    for index, job in enumerate(jobs):
        desc = normalized(job["description"])
        title_company = (normalized(job["company"]), normalized(job["title"]))
        # 空描述、空企业或空标题不能成为把不相干岗位连接起来的桥。
        for key, registry in ((desc, seen_description), (title_company if all(title_company) else None, seen_company_title)):
            if not key:
                continue
            if key in registry:
                union(index, registry[key])
            else:
                registry[key] = index
    components = defaultdict(list)
    for index, job in enumerate(jobs):
        components[find(index)].append(job["job_id"])
    families = {key: value_hash(sorted(ids))[:20] for key, ids in components.items()}
    ordered = sorted(families.values())
    random.Random(seed).shuffle(ordered)
    n_train = int(len(ordered) * ratios[0])
    n_dev = int(len(ordered) * ratios[1])
    allocation = {family: ("train" if index < n_train else "dev" if index < n_train+n_dev else "test") for index, family in enumerate(ordered)}
    for index, job in enumerate(jobs):
        job["job_family_id"] = families[find(index)]
        job["split"] = allocation[job["job_family_id"]]
    return {"family_count": len(ordered), "multi_record_families": sum(len(value) > 1 for value in components.values()),
            "largest_family": max(map(len, components.values()), default=0), "family_split_counts": dict(Counter(allocation.values())),
            "record_split_counts": dict(Counter(job["split"] for job in jobs)), "split_unit": "job_family_id",
            "split_ratio_note": "70/15/15作用于家族数，向下取整训练/开发，余数给测试；记录数占比受家族大小影响。"}


def checked_evidence(job: dict, evidence: dict) -> dict:
    field = evidence.get("field")
    start, end = evidence.get("start"), evidence.get("end")
    quote = evidence.get("quote")
    if field not in FIELDS or not isinstance(start, int) or not isinstance(end, int):
        raise ValueError("引用缺少合法字段与字符偏移。")
    text = job[field]
    if start < 0 or end <= start or end > len(text) or text[start:end] != quote:
        raise ValueError("机器引用与岗位原文不一致，拒绝输出错误证据。")
    return {"job_id": job["job_id"], "field": field, "start": start, "end": end, "quote": quote,
            "source_record_ids": job["source_record_ids"], "source_column": job["field_sources"].get(field)}


def make_edge(source_id: str, target_id: str, source_type: str, target_type: str, relation: str,
              status: str, split: str, job_id: str | None, evidence: dict | None, **extra) -> dict:
    edge = {"source_id": source_id, "target_id": target_id, "source_type": source_type, "target_type": target_type,
            "relation": relation, "status": status, "split": split, "job_id": job_id,
            "evidence": evidence, "fact_citable": status != "hypothesis", **extra}
    return {"edge_id": value_hash(edge)[:20], **edge}


def build_edges(jobs: list[dict], minimum_support: int, top_k: int) -> tuple[list[dict], dict]:
    edges = []
    category_parents = set()
    train_skills = []
    for job in jobs:
        job_id, split = job["job_id"], job["split"]
        for skill, evidence in sorted(job["skills"].items()):
            citation = checked_evidence(job, evidence)
            edges.append(make_edge(job_id, f"skill:{skill}", "job", "skill", "mentions", "observed", split, job_id, citation,
                                   skill=skill, machine_level=evidence["level"], machine_preferred=evidence["preferred"],
                                   interpretation="原文提及，并不代表必须掌握"))
        for ordinal, group in enumerate(job["groups"]):
            citation = checked_evidence(job, group["evidence"])
            if group["kind"] == "single":
                for skill in group["skills"]:
                    edges.append(make_edge(job_id, f"skill:{skill}", "job", "skill", "single_requirement", "machine_parsed", split, job_id, citation,
                                           group_kind="single", preferred=group["preferred"], scope="single"))
                continue
            group_id = "requirement_group:" + value_hash([job_id, ordinal, group["kind"], group["skills"], citation["field"], citation["start"], citation["end"]])[:20]
            # 保留逻辑组，不把 Python 或 Java 分裂成两条独立硬性要求。
            edges.append(make_edge(job_id, group_id, "job", "requirement_group", "has_requirement_group", "machine_parsed", split, job_id, citation,
                                   group_kind=group["kind"], preferred=group["preferred"], scope="any" if group["kind"] == "any" else "single"))
            for skill in group["skills"]:
                edges.append(make_edge(group_id, f"skill:{skill}", "requirement_group", "skill", "option" if group["kind"] == "any" else "member",
                                       "machine_parsed", split, job_id, citation, group_kind=group["kind"], preferred=group["preferred"]))
        if job["category"]:
            citation = checked_evidence(job, {"field":"category", "start":0, "end":len(job["category"]), "quote":job["category"]})
            edges.append(make_edge(job_id, "category:"+job["category"], "job", "category", "in_category", "observed", split, job_id, citation))
        if job["category"] and job["category_parent"]:
            # 每一折单列分类映射及其本折来源，不能由测试记录给训练图补边。
            identity = (split, job["category"], job["category_parent"])
            if identity not in category_parents:
                category_parents.add(identity)
                citation = checked_evidence(job, {"field":"category_parent", "start":0, "end":len(job["category_parent"]), "quote":job["category_parent"]})
                edges.append(make_edge("category:"+job["category"], "category_parent:"+job["category_parent"], "category", "category_parent", "category_parent", "observed", split, job_id, citation))
        if split == "train":
            train_skills.append({skill for skill, item in job["skills"].items() if item["level"] != "否定"})
    support = Counter(skill for skills in train_skills for skill in skills)
    pairs = Counter(pair for skills in train_skills for pair in combinations(sorted(skills), 2))
    neighbors = defaultdict(list)
    for (left, right), count in pairs.items():
        if count < minimum_support:
            continue
        ppmi = max(0.0, math.log(count * len(train_skills) / (support[left]*support[right])))
        neighbors[left].append((right, count, ppmi))
        neighbors[right].append((left, count, ppmi))
    for skill in sorted(neighbors):
        selected = sorted(neighbors[skill], key=lambda value: (-value[2], -value[1], value[0]))[:top_k]
        for other, count, ppmi in selected:
            edges.append(make_edge(f"skill:{skill}", f"skill:{other}", "skill", "skill", "cooccurs_hypothesis", "hypothesis", "train", None, None,
                                   support_count=count, source_support=support[skill], target_support=support[other], ppmi=round(ppmi, 8),
                                   derived_from_split="train", train_job_count=len(train_skills), not_for_fact_citation=True))
    edges.sort(key=lambda edge: edge["edge_id"])
    if len({edge["edge_id"] for edge in edges}) != len(edges):
        raise ValueError("出现重复边ID，拒绝发布不明确的图。")
    return edges, {"edge_count": len(edges), "edge_status_counts": dict(Counter(edge["status"] for edge in edges)),
                   "edge_relation_counts": dict(Counter(edge["relation"] for edge in edges)),
                   "cooccurrence_train_jobs": len(train_skills), "cooccurrence_pairs_before_filter": len(pairs),
                   "cooccurrence_min_support": minimum_support, "cooccurrence_top_k_outgoing": top_k,
                   "cooccurrence_note": "仅训练岗位贡献统计；top-k约束每个源技能的出边，假设边无原文引用。"}


def annotation_tasks(jobs: list[dict], counts: tuple[int, int, int], seed: int) -> list[dict]:
    """任务继承岗位家族主划分；label始终为null，未作任何人工判断。"""
    buckets = {split: defaultdict(list) for split in SPLITS}
    for job in jobs:
        for field in ("requirements", "description"):
            text = job[field]
            for match in re.finditer(r"[^。！？；;\n\r]+[。！？；;]?", text):
                raw = match.group()
                start = match.start()+len(raw)-len(raw.lstrip())
                end = match.end()-(len(raw)-len(raw.rstrip()))
                if end-start < 15:
                    continue
                end = min(end, start+500)
                quote = text[start:end]
                if re.search(r"不要求|无需|不必|不熟悉|未掌握|未使用|尚未|不会", quote):
                    stratum = "negation_candidate"
                elif any(g["kind"] == "any" and g["evidence"]["field"] == field and start <= g["evidence"]["start"] < end for g in job["groups"]):
                    stratum = "or_candidate"
                elif re.search(r"优先|加分|更佳|可放宽", quote):
                    stratum = "preferred_candidate"
                elif not any(e["field"] == field and start <= e["start"] < end for e in job["skills"].values()):
                    stratum = "no_dictionary_skill_candidate"
                else:
                    stratum = "other_candidate"
                task_id = value_hash([job["job_id"], field, start, end])[:20]
                task = {"task_id":task_id, "task_type":"requirement_extraction", "job_id":job["job_id"],
                        "job_family_id":job["job_family_id"], "split":job["split"], "snapshot":job["snapshot"],
                        "evidence":checked_evidence(job, {"field":field, "start":start, "end":end, "quote":quote}),
                        "text":quote, "sampling_stratum":stratum, "label":None, "annotation_status":"unlabeled",
                        "annotation_schema":{"skills":"概念及原文跨度", "logic":"single/any/all/unknown", "modality":"required/preferred/negated/unknown", "evidence":"原字段字符跨度"}}
                buckets[job["split"]][stratum].append(task)
    result = []
    for split, target in zip(SPLITS, counts):
        queues = {}
        for stratum, tasks in sorted(buckets[split].items()):
            tasks.sort(key=lambda row: value_hash([seed, row["task_id"]]))
            queues[stratum] = deque(tasks)
        selected, seen_jobs = [], set()
        while len(selected) < target and any(queues.values()):
            for stratum in sorted(queues):
                while queues[stratum] and queues[stratum][0]["job_id"] in seen_jobs:
                    queues[stratum].popleft()
                if queues[stratum] and len(selected) < target:
                    task = queues[stratum].popleft()
                    seen_jobs.add(task["job_id"])
                    selected.append(task)
        if len(selected) != target:
            raise ValueError(f"{split}可用独立岗位片段不足：需要{target}，只有{len(selected)}；未输出伪补任务。")
        result.extend(selected)
    return result


def validate_artifacts(jobs: list[dict], edges: list[dict], tasks: list[dict], raw_count: int) -> dict:
    by_id = {job["job_id"]:job for job in jobs}
    families = defaultdict(set)
    lineage = []
    for job in jobs:
        families[job["job_family_id"]].add(job["split"])
        lineage.extend(job["source_record_ids"])
    if len(lineage) != raw_count or len(set(lineage)) != raw_count:
        raise ValueError("原记录血缘没有做到一条原记录恰好归属一个版本。")
    if any(len(splits) != 1 for splits in families.values()):
        raise ValueError("岗位家族跨训练/开发/测试泄漏。")
    for edge in edges:
        if edge["status"] == "hypothesis":
            if edge["split"] != "train" or edge["evidence"] is not None or edge["fact_citable"]:
                raise ValueError("统计假设边混入事实引用或非训练统计。")
        else:
            job = by_id[edge["job_id"]]
            checked_evidence(job, edge["evidence"])
            if edge["split"] != job["split"]:
                raise ValueError("原文边划分与岗位不一致。")
    for task in tasks:
        job = by_id[task["job_id"]]
        checked_evidence(job, task["evidence"])
        if task["label"] is not None or task["split"] != job["split"] or task["job_family_id"] != job["job_family_id"]:
            raise ValueError("待标注任务伪造标签或跨主划分。")
    return {"source_lineage_complete":True, "family_split_overlap":0, "all_observed_evidence_valid":True,
            "hypothesis_never_fact_citable":True, "annotation_labels_all_null":True,
            "annotation_split_counts":dict(Counter(task["split"] for task in tasks))}


def write_json(path: Path, value: Any) -> None:
    with path.open("x", encoding="utf-8", newline="\n") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    path.chmod(0o600)


def write_jsonl(path: Path, rows: list[dict]) -> None:
    with path.open("x", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(canonical(row)+"\n")
        stream.flush()
        os.fsync(stream.fileno())
    path.chmod(0o600)


def compatible_run(directory: Path, config_hash: str) -> dict:
    if directory.is_symlink() or not directory.is_dir():
        raise ValueError("已有run不是可信的普通目录。")
    manifest_path = directory / "manifest.json"
    if not manifest_path.is_file() or manifest_path.is_symlink():
        raise ValueError("已有run无合法manifest，拒绝覆盖。")
    manifest = json.loads(manifest_path.read_text("utf-8"))
    if manifest.get("schema_version") != SCHEMA_VERSION or manifest.get("config_hash") != config_hash:
        raise ValueError("已有run配置不相容，拒绝覆盖。")
    for name, signature in manifest.get("files", {}).items():
        path = directory / name
        if Path(name).name != name or path.is_symlink() or not path.is_file() or file_hash(path) != signature["sha256"]:
            raise ValueError("已有run文件校验失败，拒绝覆盖损坏数据。")
    if set(manifest.get("files", {})) != {"jobs.jsonl", "edges.jsonl", "annotation_tasks.jsonl"}:
        raise ValueError("已有run文件清单不完整。")
    return manifest


def build_dataset(source: Path, output_root: Path, repository: Path, seed: int = 42,
                  ratios: tuple[float,float,float] = (.70,.15,.15), annotation_counts: tuple[int,int,int] = (300,100,200),
                  minimum_support: int = 10, top_k: int = 10) -> tuple[Path, dict]:
    if not source.is_absolute() or not output_root.is_absolute() or not repository.is_absolute():
        raise ValueError("输入、输出根目录与仓库路径均须为绝对路径。")
    source, output_root, repository = source.resolve(), output_root.resolve(), repository.resolve()
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        raise ValueError("输入必须为现有xlsx文件。")
    if output_root.is_relative_to(repository):
        raise ValueError("研究输出必须在仓库外，禁止将真实岗位与公司文本写回仓库。")
    if len(ratios) != 3 or any(value < 0 for value in ratios) or not math.isclose(sum(ratios), 1.0):
        raise ValueError("划分比例必须是和为1的三个非负数。")
    if len(annotation_counts) != 3 or any(count < 0 for count in annotation_counts) or minimum_support < 1 or top_k < 1:
        raise ValueError("任务数非负，统计支持阈值和top-k至少为1。")
    Job = load_domain(repository)
    snapshot = file_hash(source)
    config = {"schema_version":SCHEMA_VERSION, "snapshot":snapshot, "seed":seed, "split_ratios":list(ratios),
              "annotation_counts":list(annotation_counts), "cooccurrence_min_support":minimum_support, "cooccurrence_top_k":top_k,
              "code_sha256":file_hash(Path(__file__)), "domain_sha256":file_hash(repository/"job_agent/domain.py"),
              "skill_rules_sha256":file_hash(repository/"scripts/profile_data.py")}
    config_hash = value_hash(config)
    run_name = f"snapshot-{snapshot[:12]}-seed{seed}-{config_hash[:12]}"
    run_dir = output_root/run_name
    output_root.mkdir(parents=True, exist_ok=True, mode=0o700)
    if run_dir.exists() or run_dir.is_symlink():
        return run_dir, compatible_run(run_dir, config_hash)
    lock = output_root / ("."+run_name+".lock")
    try:
        fd = os.open(lock, os.O_CREAT|os.O_EXCL|os.O_WRONLY, 0o600)
    except FileExistsError as exc:
        raise ValueError("同一run正在构建或有待检查的旧锁，不会覆盖。") from exc
    os.close(fd)
    staging = Path(tempfile.mkdtemp(prefix="."+run_name+"-", dir=output_root))
    try:
        jobs, read_stats = read_jobs(source, Job, snapshot)
        if not jobs:
            raise ValueError("输入没有岗位记录。")
        split_stats = assign_families(jobs, seed, ratios)
        edges, graph_stats = build_edges(jobs, minimum_support, top_k)
        tasks = annotation_tasks(jobs, annotation_counts, seed)
        checks = validate_artifacts(jobs, edges, tasks, read_stats["raw_records"])
        if file_hash(source) != snapshot:
            raise RuntimeError("读取期间输入快照发生变化，拒绝发布。")
        files = {}
        for name, rows in (("jobs.jsonl",jobs), ("edges.jsonl",edges), ("annotation_tasks.jsonl",tasks)):
            write_jsonl(staging/name, rows)
            files[name] = {"sha256":file_hash(staging/name), "rows":len(rows)}
        manifest = {"schema_version":SCHEMA_VERSION, "config_hash":config_hash, "config":config,
                    "created_utc":datetime.now(timezone.utc).isoformat(timespec="seconds"), "source_path":str(source), "source_sha256":snapshot,
                    "source_unchanged":True, "source_read_only":True, "data":read_stats, "splits":split_stats, "graph":graph_stats,
                    "checks":checks, "files":files, "annotation_note":"任务标签全部为空，尚未进行人工标注。",
                    "family_note":"同规范职责或同企业规范标题的连通种子簇，不是已验证语义家族。",
                    "privacy_note":"包含岗位原文与企业，仅存仓库外私有研究目录，不应直接提交公开仓库。"}
        write_json(staging/"manifest.json", manifest)
        if file_hash(source) != snapshot:
            raise RuntimeError("发布前输入快照发生变化，拒绝发布。")
        if run_dir.exists() or run_dir.is_symlink():
            raise ValueError("发布位置已存在，拒绝覆盖不明确的run。")
        os.rename(staging, run_dir)
        if file_hash(source) != snapshot:
            raise RuntimeError("发布后源文件发生外部变化，完整性检查未通过。")
        return run_dir, manifest
    finally:
        if staging.exists():
            shutil.rmtree(staging)
        lock.unlink(missing_ok=True)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-root", required=True, type=Path)
    parser.add_argument("--repo-root", type=Path, default=Path.cwd())
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--split-ratios", type=float, nargs=3, default=(.70,.15,.15), metavar=("训练","开发","测试"))
    parser.add_argument("--annotation-counts", type=int, nargs=3, default=(300,100,200), metavar=("训练","开发","测试"))
    parser.add_argument("--minimum-support", type=int, default=10)
    parser.add_argument("--top-k", type=int, default=10)
    args = parser.parse_args(argv)
    try:
        directory, manifest = build_dataset(args.input, args.output_root, args.repo_root, args.seed, tuple(args.split_ratios), tuple(args.annotation_counts), args.minimum_support, args.top_k)
    except (ValueError, RuntimeError, OSError) as exc:
        parser.exit(2, f"研究数据构建失败：{exc}\n")
    # stdout仅含目录和聚合，不打印企业、岗位原文、行级ID或标注片段。
    print(json.dumps({"run_dir":str(directory), "schema_version":manifest["schema_version"], "source_sha256":manifest["source_sha256"],
                      "data":manifest["data"], "splits":manifest["splits"], "graph":manifest["graph"], "checks":manifest["checks"]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
