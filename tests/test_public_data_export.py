"""公开导出的个人信息清理、企业稳定性、文件一致性与原表保全。"""
import csv
import gzip
import hashlib
import json
from collections import Counter

from openpyxl import Workbook, load_workbook
import pytest

from scripts.export_public_jobs import ARTIFACTS, FIELDS, HEADERS, coarse_address, export_public_jobs, sanitize_text


def test_contact_redaction_keeps_skills_and_salary():
    text = "熟悉Python、SQL和微信小程序。联系人：测试先生；电话：１３８００００００００；微信：demo_id。使用Vue开发页面。"
    result = sanitize_text(text, "样例企业", Counter())
    for forbidden in ["138", "测试先生", "demo_id"]:
        assert forbidden not in result
    for required in ["Python", "SQL", "微信小程序", "Vue"]:
        assert required in result
    for text in ["138 0000 0000", "010-12345678", "demo@example.test", "https://example.test/p", "123456789012345678"]:
        assert sanitize_text(text, "样例企业", Counter()) == "[联系信息已移除]"
    assert sanitize_text("15-25K·14薪", "样例企业", Counter()) == "15-25K·14薪"
    assert sanitize_text('=HYPERLINK("x")', "样例企业", Counter()).startswith("'=")
    for text in ["demo (at) example [dot] test", "example.com", "薇信demo_abc", "QQ群12345678", "一三八零零零零零零零零"]:
        assert sanitize_text(text, "样例企业", Counter()) == "[联系信息已移除]"
    assert sanitize_text("熟悉ASP.NET和Socket.IO", "样例企业", Counter()) == "熟悉ASP.NET和Socket.IO"


def test_address_preserves_unknown_city():
    assert coarse_address("深圳市南山区测试路100号9层") == "深圳市南山区"
    assert coarse_address("南山区测试路100号") == "南山区"
    assert coarse_address("广东某科技园8楼") == ""
    assert coarse_address("北京市海淀区测试路1号") == "北京市海淀区"


def test_exports_keep_source_and_versions_and_formats_equal(tmp_path):
    source = tmp_path / "original.xlsx"
    book = Workbook(); sheet = book.active; sheet.append(FIELDS)
    for company, description in [("样例甲", "使用Python开发接口。"), ("样例甲", "使用Python清洗数据。"), ("样例乙", "邮箱demo@example.test；熟悉SQL。")]:
        sheet.append(["开发", company, "10-20K", "本科", description, "深圳南山区1栋5层", "开发", "技术"])
    book.save(source); book.close()
    before = source.read_bytes(); output = tmp_path / "public"
    manifest = export_public_jobs(source, output)
    assert source.read_bytes() == before
    assert manifest["rows"] == 3 and manifest["company_count"] == 2 and manifest["business_key_unique"] == 2
    with gzip.open(output / ARTIFACTS[1], "rt", encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.DictReader(stream))
    assert csv_rows[0]["企业"] == csv_rows[1]["企业"] != csv_rows[2]["企业"]
    assert len({row["记录编号"] for row in csv_rows}) == 3
    book = load_workbook(output / ARTIFACTS[0], read_only=True)
    rows = list(book.active.values); book.close()
    assert list(rows[0]) == HEADERS
    xlsx_rows = [{key: value or "" for key, value in zip(HEADERS, row)} for row in rows[1:]]
    assert xlsx_rows == csv_rows
    for name, data in manifest["files"].items():
        assert hashlib.sha256((output / name).read_bytes()).hexdigest() == data["sha256"]
    assert not any(manifest["residual_pattern_cells"].values())
    assert "demo@example.test" not in json.dumps(csv_rows)
    from scripts.verify_public_data import verify
    assert verify(output)["rows"] == 3
    with pytest.raises(FileExistsError): export_public_jobs(source, output)
    with (output / ARTIFACTS[1]).open("ab") as stream: stream.write(b"tampered")
    with pytest.raises(ValueError, match="摘要不匹配"): verify(output)


def test_unreviewed_columns_are_rejected(tmp_path):
    source = tmp_path / "extra.xlsx"
    book = Workbook(); sheet = book.active; sheet.append([*FIELDS, "联系人"])
    sheet.append(["岗位", "企业", "", "", "", "", "", "", "待审核人员"])
    book.save(source); book.close()
    with pytest.raises(ValueError, match="8列模式"):
        export_public_jobs(source, tmp_path / "public")
    assert not (tmp_path / "public").exists()
