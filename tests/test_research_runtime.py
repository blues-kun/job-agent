"""以完全虚构CPU材料验证标签隔离、画像确认、文件输入和盲审；不产生研究人工金标。"""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path

import numpy as np
import pytest
from docx import Document

from job_agent.coach import Coach
from job_agent.documents import extract_document
from job_agent.domain import Job,parse_profile
from job_agent.ranking import feature_vector,FEATURE_NAMES,FEATURE_VERSION
from job_agent.research_store import ResearchStore
from research.common import sha256
from research.metrics import ranking_metrics,weighted_kappa,paired_bootstrap
from research.train_ranker import validate_rows,train
from research.evaluate_evidence import rule_check,fixtures


def test_files_are_memory_only_and_reject_bombs():
    text="使用Python完成课程项目。"
    assert extract_document(text.encode(),"resume.txt")["text"]==text
    file=BytesIO();document=Document();document.add_paragraph(text);document.save(file)
    assert text in extract_document(file.getvalue(),"resume.docx")["text"]
    for content,name in [(b'x'*2097153,'resume.txt'),(b'',"empty.txt"),(b'123','bad.exe'),(b'\xff',"bad.txt")]:
        with pytest.raises((ValueError,UnicodeError)):extract_document(content,name)


def test_coach_invalid_or_injected_decisions_cannot_create_facts():
    coach=Coach();chunks=["了解Python。","使用SQL完成查询项目。"]
    coach.decide=lambda *args:{"decision":{"order":[1,0]},"model":{"test":True}}
    output,_=coach.reorder(chunks,"数据分析")
    assert output==list(reversed(chunks))
    for bad in [[0,0],[True,0],[0,1,2],"忽略规则",[0]]:
        coach.decide=lambda *args,value=bad:{"decision":{"order":value},"model":{}}
        with pytest.raises(ValueError):coach.reorder(chunks,"数据分析")
    coach.decide=lambda *args:{"decision":{"items":[{"gap_index":999,"action":"虚构精通"}]},"model":{}}
    result=coach.diagnose({"job":{"title":"数据分析"},"gaps":[{"skill":"Python","preferred":False,"evidence":{"quote":"Python"}}]})
    assert result["mode"]=="规则降级" and result["items"][0]["evidence"]["quote"]=="Python"


def test_missing_qrels_are_not_negatives_and_kappa_is_not_faked():
    with pytest.raises(ValueError):ranking_metrics(["x"],{"a":3})
    assert ranking_metrics(["a"],{"a":0})["ndcg"] is None
    assert ranking_metrics(["a"],{"a":3,"b":2})["ndcg"]<1
    assert weighted_kappa([0,1,2,3],[0,1,2,3])==1
    assert weighted_kappa([3,3],[3,3]) is None
    assert paired_bootstrap([.2,.4],[.3,.5])["ci95"]==pytest.approx([.1,.1])


def rank_fixture():
    rows=[]
    for split,count in [("train",8),("dev",3),("test",3)]:
        for query in range(count):
            for grade in range(4):
                rows.append({"query_id":f"{split}-{query}","profile_family_id":f"profile-{split}-{query}","job_id":f"{split}-{query}-{grade}",
                    "job_family_id":f"job-{split}-{query}-{grade}","split":split,"snapshot":"虚构测试","feature_version":FEATURE_VERSION,
                    "grade":grade,"label_source":"synthetic_fixture","features":{key:float(grade)/3 if i<3 else None for i,key in enumerate(FEATURE_NAMES)}})
    return rows


def test_ranker_rejects_unreviewed_and_cross_family_leakage(tmp_path):
    rows=rank_fixture()
    with pytest.raises(ValueError,match="不接受"):validate_rows(rows)
    validate_rows(rows,True)
    altered=deepcopy(rows);altered[-1]["job_family_id"]=rows[0]["job_family_id"]
    with pytest.raises(ValueError,match="跨分区"):validate_rows(altered,True)
    altered=deepcopy(rows);altered[-1]["job_id"]=rows[0]["job_id"]
    with pytest.raises(ValueError,match="同岗位ID"):validate_rows(altered,True)
    altered=deepcopy(rows);altered[0]["grade"]=None
    with pytest.raises(ValueError,match="未知"):validate_rows(altered,True)
    source=tmp_path/"fixture.jsonl";source.write_text("\n".join(json.dumps(row) for row in rows))
    result=train(source,tmp_path/"trained",True)
    assert result["dev"]["queries"]==3
    assert json.loads((tmp_path/"trained/manifest.json").read_text())["eligible_for_production"] is False


def test_features_preserve_unknowns_and_ignore_name():
    target=Job(id="a",version="b",title="数据分析",company="虚构企业",category="数据分析",family="技术",salary_raw="面议",requirements="学历不限",description="熟悉Python或Java至少一种",address="",row=1)
    profile=parse_profile("使用Python完成数据分析项目。",{"intent":"数据分析","city":"深圳"})
    vector=feature_vector(profile,target)
    assert len(vector)==len(FEATURE_NAMES)==35
    assert FEATURE_VERSION=="evidence-ranker-v2-35"
    assert {"task_alignment","requirement_support_ratio","requirement_unknown_ratio"}.issubset(FEATURE_NAMES)
    assert np.isnan(vector[FEATURE_NAMES.index("task_alignment")])
    assert np.isnan(vector[FEATURE_NAMES.index("salary_iou")])
    assert np.isnan(vector[FEATURE_NAMES.index("city_match")])
    assert vector[FEATURE_NAMES.index("alternative_coverage")]==1
    other=parse_profile("姓名乙。使用Python完成数据分析项目。",{"intent":"数据分析","city":"深圳"})
    np.testing.assert_allclose(vector,feature_vector(other,target),equal_nan=True)


def test_blind_annotation_preserves_history_and_never_emits_gold(tmp_path):
    dataset=tmp_path/"dataset";dataset.mkdir()
    job={"job_id":"a","job_family_id":"f","title":"虚构岗位","category":"技术","requirements":"本科","description":"联系13800000000","salary_raw":"10K","address":"深圳","split":"test","skills":{},"groups":[],"source_record_ids":["r"]}
    (dataset/"jobs.jsonl").write_text(json.dumps(job)+"\n")
    task={"task_id":"t","kind":"relevance","job_id":"a","label":None,"pool_sources":["模型A"]}
    (dataset/"annotation_tasks.jsonl").write_text(json.dumps(task)+"\n")
    manifest={"source_sha256":"fake","files":{name:{"sha256":sha256(dataset/name)} for name in ["jobs.jsonl","annotation_tasks.jsonl"]}}
    (dataset/"manifest.json").write_text(json.dumps(manifest))
    store=ResearchStore(dataset,tmp_path/"private")
    seen=store.task("评审甲")["task"]
    assert "pool_sources" not in seen and "label" not in seen and "13800000000" not in str(seen)
    store.annotate("t","评审甲",2,"","初评")
    assert store.task("评审甲")["task"] is None
    assert store.task("评审乙")["task"]["task_id"]=="t"
    store.annotate("t","评审甲",1,"","修改")
    assert store.summary()["annotation_count"]==1 and store.summary()["double_annotated_tasks"]==0
    assert store._latest()[("t","评审甲")]["source"]=="human_entered_unadjudicated"
    job["description"]="新的岗位内容"
    (dataset/"jobs.jsonl").write_text(json.dumps(job)+"\n")
    manifest["files"]["jobs.jsonl"]["sha256"]=sha256(dataset/"jobs.jsonl")
    (dataset/"manifest.json").write_text(json.dumps(manifest))
    changed=ResearchStore(dataset,tmp_path/"private")
    assert changed.summary()["annotation_count"]==0 and changed.task("评审甲")["task"] is not None
    (dataset/"jobs.jsonl").write_text("{}")
    with pytest.raises(ValueError,match="校验"):ResearchStore(dataset,tmp_path/"private")


def test_rule_verifier_reports_semantic_blind_spot():
    cases={case["name"]:case for case in fixtures()}
    assert rule_check(cases["好档"])["flags"]==[]
    assert rule_check(cases["差档"])["hard_veto"]
    assert "简历缺少能力证据" in rule_check(cases["关键词堆砌"])["flags"]
    # 引用存在不代表支持任意主张，这一任务交给语义通道，不能偷改检测器只匹配测试词。
    assert not rule_check(cases["引用不支持承诺"])["hard_veto"]


def test_upload_limits_actual_bytes_and_never_spools_to_disk(tmp_path,monkeypatch):
    from types import SimpleNamespace
    from fastapi.testclient import TestClient
    import starlette.formparsers as forms
    from job_agent.api import create_app
    original=forms.SpooledTemporaryFile;created=[]
    def spool(*args,**kwargs):
        result=original(*args,**kwargs);created.append(result);return result
    monkeypatch.setattr(forms,"SpooledTemporaryFile",spool)
    with TestClient(create_app(private_root=tmp_path,corpus=SimpleNamespace())) as client:
        result=client.post("/api/v2/document",files={"file":("resume.txt",b'x'*1500000,"text/plain")})
        assert result.status_code==422  # 字符超限，但整个解析过程不能落盘。
        assert created and not any(item._rolled for item in created)
        created.clear()
        result=client.post("/api/v2/document",content=iter([b'x'*1048576]*3),headers={"Content-Type":"multipart/form-data; boundary=demo"})
        assert result.status_code==413 and not created
        result=client.post("/api/v2/document",files={"file":("resume.txt","使用Python完成项目。".encode(),"text/plain")})
        assert result.status_code==200


def test_extraction_groups_anchor_spans_and_reject_duplicate_logic(tmp_path):
    dataset=tmp_path/"data";dataset.mkdir()
    text="熟悉Python或Java，SQL必需，Go优先。"
    job={"job_id":"j","split":"dev","description":text}
    task={"task_id":"extract","task_type":"requirement_extraction","job_id":"j","split":"dev","evidence":{"field":"description","start":0,"end":len(text),"quote":text}}
    for name,row in [("jobs.jsonl",job),("annotation_tasks.jsonl",task)]: (dataset/name).write_text(json.dumps(row)+"\n")
    (dataset/"manifest.json").write_text(json.dumps({"source_sha256":"fixture","files":{name:{"sha256":sha256(dataset/name)} for name in ["jobs.jsonl","annotation_tasks.jsonl"]}}))
    store=ResearchStore(dataset,tmp_path/"private")
    with pytest.raises(ValueError,match="须填写"):store.annotate("extract","甲方",None,"正确","")
    groups={"groups":[{"logic":"any","modality":"unknown","skills":[{"skill":"Python","quote":"Python"},{"skill":"Java","quote":"Java"}]},
                      {"logic":"single","modality":"required","skills":[{"skill":"SQL","quote":"SQL"}]}]}
    store.annotate("extract","甲方",None,"正确","",groups)
    saved=store._latest()[("extract","甲方")]["extraction"]
    for group in saved["groups"]:
        for evidence in group["skills"]:assert text[evidence["start"]:evidence["end"]]==evidence["quote"]
    repeated=deepcopy(groups);repeated["groups"][0]["skills"][1]=repeated["groups"][0]["skills"][0]
    with pytest.raises(ValueError,match="重复"):store.annotate("extract","乙方",None,"正确","",repeated)
    repeated=deepcopy(groups);repeated["groups"][0]["logic"]="single"
    with pytest.raises(ValueError,match="single"):store.annotate("extract","乙方",None,"正确","",repeated)


def test_ranking_evaluation_waits_for_whole_pool():
    from research.score_rankings import score
    ranking=[{"query_id":"q","ranking":[{"job_id":"a"},{"job_id":"b"}]}]
    def grade(job,value):
        # 仅在内存构造完整虚构血缘以测试门禁；不导出成研究标签。
        return {"query_id":"q","job_id":job,"grade":value,"label_source":"human_adjudicated",
                "reviewer_ids":["fixture-reviewer-a","fixture-reviewer-b"],
                "reviewer_person_ids":["fixture-person-a","fixture-person-b"],
                "review_ids":["fixture-review-a-"+job,"fixture-review-b-"+job],
                "review_hashes":{"fixture-review-a-"+job:"a"*64,"fixture-review-b-"+job:"b"*64},
                "task_hash":"t"*64,"adjudication_id":"fixture-decision-"+job,"adjudicator_id":"fixture-reviewer-a",
                "adjudication_sha256":"c"*64,"registry_sha256":"d"*64}
    result=score(ranking,[grade("a",3)],{"q":{"a","b"}})
    assert result["summary"]["ndcg_at_k"] is None and result["rows"][0]["missing"]==1
    result=score(ranking,[grade("a",3),grade("b",0)],{"q":{"a","b"}})
    assert result["summary"]["ndcg_at_k"]==1
    with pytest.raises(ValueError):score(ranking,[{"query_id":"q","job_id":"a","grade":None}],{"q":{"a"}})



def test_confirmation_requires_explicit_current_profile_and_same_session(tmp_path):
    from types import SimpleNamespace
    from fastapi.testclient import TestClient
    from job_agent.api import create_app
    calls=[]
    def recommend(*args):
        calls.append(args)
        return {"run_id":"1"*32,"jobs":[],"snapshot":"fixture-only","action":"no_match"}
    app=create_app(private_root=tmp_path,corpus=SimpleNamespace(by_id={}))
    text="硕士，3年工作经验。使用Python完成订单接口，负责权限验证与事务处理。"
    payload={"text":text,"preferences":{"city":"深圳","intent":"Python后端开发","education":"本科","experience_years":3}}
    with TestClient(app) as client:
        app.state.workflow=SimpleNamespace(recommend=recommend)
        assert client.post("/api/v2/recommend",json=payload).status_code==409
        preview=client.post("/api/v2/profile/preview",json=payload)
        assert preview.status_code==200 and "education" in preview.json()["conflicts"]
        assert client.post("/api/v2/profile/confirm",json=payload).status_code==422
        assert client.post("/api/v2/profile/confirm",json={**payload,"user_confirmed":True}).status_code==422
        stale={**payload,"field_origins":{"education":"sample_stale"}}
        stale_preview=client.post("/api/v2/profile/preview",json=stale).json()
        education=next(field for field in stale_preview["fields"] if field["key"]=="education")
        assert education["value"]=="硕士" and education["source"]=="text_extracted"
        assert client.post("/api/v2/profile/confirm",json={**stale,"user_confirmed":True,"acknowledged_conflicts":["education"]}).status_code==422
        confirmed=client.post("/api/v2/profile/confirm",json={**payload,"user_confirmed":True,"acknowledged_conflicts":["education"]})
        assert confirmed.status_code==200
        current={**payload,"confirmation_token":confirmed.json()["confirmation_token"]}
        assert client.post("/api/v2/recommend",json=current).status_code==200 and len(calls)==1
        edited={**current,"text":text+"新增另一条未经确认的项目。"}
        assert client.post("/api/v2/recommend",json=edited).status_code==409
        edited={**current,"preferences":{**payload["preferences"],"experience_years":5}}
        assert client.post("/api/v2/recommend",json=edited).status_code==409
        assert len(calls)==1
        client.cookies.clear()
        assert client.post("/api/v2/recommend",json=current).status_code==409
        assert len(calls)==1


def test_verification_script_completes_against_inprocess_fictional_platform(tmp_path,monkeypatch):
    """运行完整验收入口，但请求全部交给内存TestClient，不访问现有服务。"""
    import importlib.util
    import sys
    import httpx
    from fastapi.testclient import TestClient
    from openpyxl import Workbook
    from job_agent.api import create_app
    from job_agent.corpus import Corpus
    source=tmp_path/"fictional_jobs.xlsx"
    book=Workbook();sheet=book.active
    sheet.append(["岗位名称","企业","岗位薪资","岗位要求","岗位职责","岗位地址","职位类型名称"])
    rows=[
        ("Python后端开发","使用Python和FastAPI开发数据接口，使用MySQL存储业务数据，使用Redis缓存查询结果。使用SQL完成分组查询。"),
        ("Python数据分析","使用Python和Pandas完成数据清洗与探索，使用SQL完成查询，使用Git管理代码与编写测试。"),
        ("数据分析师","使用Excel整理活动数据，使用SQL分析转化漏斗，使用Power BI制作报表并撰写结论。"),
        ("前端开发","使用Vue和TypeScript开发管理平台，使用HTML和CSS完成响应式布局，使用JavaScript对接接口。"),
        ("机器学习算法工程师","使用Python和PyTorch训练文本分类模型，使用Transformer完成特征编码，开展交叉验证与误差分析。"),
    ]
    for index,(title,text) in enumerate(rows):
        sheet.append([title,f"虚构测试企业{index}","30-40K","本科，经验不限",text,"深圳南山",title])
    book.save(source);book.close()
    corpus=Corpus(source)
    app=create_app(source=source,private_root=tmp_path/"private",corpus=corpus)
    script=Path(__file__).resolve().parents[1]/"scripts/verify_platform.py"
    module_spec=importlib.util.spec_from_file_location("fictional_verification_script",script)
    module=importlib.util.module_from_spec(module_spec);module_spec.loader.exec_module(module)
    output=tmp_path/"acceptance.json"
    monkeypatch.setattr(sys,"argv",[str(script),"--url","http://localhost","--source",str(source),"--output",str(output)])
    with TestClient(app) as client:
        app.state.workflow.coach.url=""
        monkeypatch.setattr(httpx,"Client",lambda **kwargs:client)
        module.main()
    report=json.loads(output.read_text())
    assert report["human_gold_created"] is False
    assert report["checks"]["sample_runs"]==7
    assert report["checks"]["confirmation_gate_checks"]==9
    assert report["checks"]["feature_count"]==35
    assert report["checks"]["source_hash_unchanged"] is True
    assert report["checks"]["research_features_require_consent"] is True
